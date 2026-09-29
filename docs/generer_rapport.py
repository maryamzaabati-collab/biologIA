#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


VERT = RGBColor(0x1F, 0x4D, 0x3A)
ENCRE = RGBColor(0x1A, 0x1F, 0x1C)
GRIS = RGBColor(0x4C, 0x58, 0x50)


def set_run(run, size=11, bold=False, italic=False, color=ENCRE, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_para(doc, text, size=11, bold=False, italic=False, space_after=8, space_before=0, align=None, color=ENCRE):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run(run, size=16 if level == 1 else 13, bold=True, color=VERT)
    # Petite ligne sous les titres de niveau 1
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pBdr = pPr.makeelement(qn("w:pBdr"), {})
        bottom = pBdr.makeelement(
            qn("w:bottom"),
            {
                qn("w:val"): "single",
                qn("w:sz"): "12",
                qn("w:space"): "4",
                qn("w:color"): "1F4D3A",
            },
        )
        pBdr.append(bottom)
        pPr.append(pBdr)
    return p


def add_bullets(doc, items, bold_lead=True):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(1.1)
        if isinstance(item, tuple):
            titre, suite = item
            r1 = p.add_run(titre)
            set_run(r1, size=11, bold=True)
            r2 = p.add_run(suite)
            set_run(r2, size=11, bold=False)
        else:
            r = p.add_run(item)
            set_run(r, size=11)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(1.1)
        if isinstance(item, tuple):
            titre, suite = item
            r1 = p.add_run(titre)
            set_run(r1, size=11, bold=True)
            r2 = p.add_run(suite)
            set_run(r2, size=11)
        else:
            r = p.add_run(item)
            set_run(r, size=11)


def add_callout(doc, titre, texte):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    r1 = p.add_run(titre + " ")
    set_run(r1, size=11, bold=True, italic=True, color=VERT)
    r2 = p.add_run(texte)
    set_run(r2, size=11, italic=True, color=GRIS)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    header = section.header
    hp = header.paragraphs[0]
    hr = hp.add_run("Registre de traçabilité — Projet 4 Biologie / Laboratoires — BUT Science des Données")
    set_run(hr, size=8, italic=True, color=GRIS)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Document pédagogique — données fictives uniquement  ·  p. ")
    set_run(fr, size=8, color=GRIS)
    from docx.oxml import OxmlElement

    def add_page_field(paragraph):
        run = paragraph.add_run()
        set_run(run, size=8, color=GRIS)
        fldChar1 = OxmlElement("w:fldChar")
        fldChar1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "end")
        run._r.append(fldChar1)
        run._r.append(instr)
        run._r.append(fldChar2)

    add_page_field(fp)

    add_para(doc, "BUT Science des Données  ·  Projet 4", size=12, color=GRIS, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(
        doc,
        "Registre de traçabilité des données de référence",
        size=22,
        bold=True,
        color=VERT,
        space_after=4,
        space_before=16,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "Rapport pédagogique — groupe et enseignants",
        size=12,
        italic=True,
        color=GRIS,
        space_after=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        "Prototype de gouvernance des lots de calibration (données fictives). "
        "Ce n’est ni un manuel technique, ni un dossier d’accréditation.",
        size=11,
        space_after=8,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_para(
        doc,
        f"{date.today().strftime('%d/%m/%Y')}  ·  http://localhost:3000  ·  "
        "samira / labo2026 (technicien)  ·  amrani / labo2026 (biologiste)",
        size=10,
        color=GRIS,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=16,
    )

    add_heading_custom(doc, "1. Contexte et objectif du projet", 1)
    add_para(
        doc,
        "Derrière un résultat de labo (glycémie, CRP, TSH…), il y a un appareil et des données de référence : "
        "des lots de contrôle qui servent à calibrer les seuils et à vérifier que la machine n’a pas dérivé. "
        "Si un lot n’a ni machine, ni date, ni conditions, ni validation biologiste — ou s’il contient encore des identités — "
        "on ne peut plus dire d’où vient la règle d’alerte. Le sujet n’est pas de diagnostiquer un patient : "
        "c’est de retracer la donnée de calibration.",
    )
    add_para(
        doc,
        "C’est important pour trois raisons. L’ISO 15189 (accréditation COFRAC) exige de retracer méthodes et références, "
        "et de les faire valider par un biologiste. Le RGPD interdit de laisser des identités dans des fichiers de travail. "
        "L’AI Act demande, pour un outil d’aide, de documenter l’origine des données qui nourrissent une règle automatique. "
        "Sans registre, tout reste dans des tableurs et une validation orale : le jour d’un audit, on ne sait plus ce qui a servi.",
    )
    add_para(doc, "Périmètre de l’application :")
    add_bullets(
        doc,
        [
            ("Machines : ", "analyseurs (nom, type)."),
            ("Lots : ", "jeu de mesures de contrôle, avec date et conditions de mesure."),
            ("Pseudonymisation : ", "identités remplacées par un code ; scan des noms de colonnes d’un CSV."),
            ("Validation : ", "le technicien enregistre ; seul un biologiste signe le lot et les comptes."),
            ("Contrôle automatique : ", "signalements si le dossier est incomplet, règles liées au lot, cartes de dérive."),
        ],
    )
    add_callout(
        doc,
        "Hors sujet, par choix :",
        "pas de dossier patient, pas de diagnostic. Les données de démo sont fictives.",
    )

    add_heading_custom(doc, "2. Concepts clés", 1)
    add_bullets(
        doc,
        [
            (
                "Lot. ",
                "Ce n’est pas un patient. C’est le dossier d’un jeu de mesures de contrôle (souvent un CSV) : "
                "machine, date, conditions, statut, score 0–100, code unique LT-…. "
                "Analogie : l’étiquette d’un réactif — on doit pouvoir remonter au flacon, pas seulement au résultat affiché.",
            ),
            (
                "Machine. ",
                "L’analyseur qui a produit les mesures. On sépare les lots par appareil pour voir si c’est une méthode qui dérive, ou un équipement.",
            ),
            (
                "Validation. ",
                "Le technicien constitue le dossier ; le biologiste signe que le lot peut servir de référence. "
                "Signature nominative, datée ; au-delà de 12 mois le lot passe « à revalider ».",
            ),
            (
                "Signalement. ",
                "Alerte automatique si le dossier est incomplet (machine, date, conditions, pseudonymisation) "
                "ou si le CSV a des colonnes identifiantes (nom, NIR…). On ne peut pas le fermer sans avoir corrigé le lot.",
            ),
        ],
    )
    add_para(doc, "Contrôle qualité automatique, en trois étages :")
    add_bullets(
        doc,
        [
            ("Dossier complet ? ", "Sinon, signalement."),
            ("Règle d’alerte. ", "Seuils (bas / haut) accrochés à un lot, donc à une origine. Un simulateur teste une valeur fictive — pas un vrai bilan."),
            ("Carte de contrôle. ", "Moyennes des lots dans le temps, par machine. Un point à plus de 2 écarts-types est hors norme (esprit Levey-Jennings, version simplifiée)."),
        ],
    )
    add_para(
        doc,
        "Rôles : client (lecture), technicien (enregistre et corrige), biologiste (valide les lots et les comptes). "
        "Une inscription reste en attente jusqu’à acceptation par un biologiste.",
    )

    add_heading_custom(doc, "3. Fonctionnalités livrées", 1)
    add_para(doc, "Ce que l’appli permet de faire aujourd’hui, et à quoi ça sert en labo :")
    add_bullets(
        doc,
        [
            ("Connexion et comptes. ", "Identifiant / mot de passe (hachés). Inscription avec rôle demandé ; un biologiste accepte ou refuse. Google est optionnel (si configuré). Sert à ce que le rôle ne soit plus un menu cosmétique."),
            ("Tableau de bord. ", "Effectifs, signalements, lots à traiter, ouverture par code LT-…. Sert à voir le retard qualité en un coup d’œil."),
            ("Machines. ", "Parc d’analyseurs. Sert à ancrer chaque lot sur un équipement."),
            ("Lots. ", "Création, édition, CSV, statuts (conforme / à vérifier / à valider / à revalider), historique des champs, score 0–100. Sert à constituer le dossier de référence à la source."),
            ("Pseudonymisation. ", "Case obligatoire + alerte si un en-tête CSV ressemble à une identité. Sert à ne pas calibrer une règle sur un fichier nominatif."),
            ("Validation biologiste. ", "Seuls les biologistes connectés peuvent signer ; le serveur refuse les autres. Sert de tampon nominatif et daté."),
            ("Signalements. ", "Liste des non-conformités ouvertes ; résolution bloquée tant que le motif est vrai. Sert de cahier qualité, pas de liste décorative."),
            ("Règles et simulateur. ", "Seuils liés à un lot ; test d’une valeur fictive. Sert à montrer qu’une alerte a une origine, sans diagnostiquer."),
            ("Cartes de contrôle. ", "Dérive des moyennes par machine (± 2 s). Sert à surveiller l’analyseur."),
            ("Certificat imprimable. ", "Fiche lot + code-barres. Sert de feuille d’audit pédagogique (esprit ISO 15189 / 17025)."),
            ("Synthèse PDF / CSV. ", "Totaux du registre. Sert de fiche pour une réunion qualité interne — pas un rapport COFRAC."),
            ("Page Conformité. ", "Lien explicite ISO 15189, RGPD, AI Act, chaîne Machine → Lot → Règle → Validation. Sert à l’oral ; le logiciel n’est pas certifié."),
        ],
    )

    add_heading_custom(doc, "4. Architecture technique", 1)
    add_para(
        doc,
        "Serveur Node.js / Express, pages HTML et JavaScript sans framework (pas de React), données dans un fichier SQLite "
        "placé hors OneDrive (dossier utilisateur) pour rester ouvrable. Après connexion, un cookie de session porte le rôle ; "
        "les droits sont vérifiés côté serveur, pas seulement à l’écran.",
    )

    add_heading_custom(doc, "5. Limites actuelles (et pourquoi)", 1)
    add_para(doc, "C’est un démonstrateur de soutenance, pas un logiciel de production.")
    add_bullets(
        doc,
        [
            (
                "Pas une certification. ",
                "La page Conformité explique ISO / RGPD / AI Act ; elle n’accrédite rien. Un projet BUT n’a ni DPO ni hébergement HDS : afficher « conforme » comme un tampon serait trompeur.",
            ),
            (
                "Scan CSV limité aux en-têtes. ",
                "Une colonne « col1 » remplie d’identités passe ; analyser le contenu des cellules serait un autre projet, et stocker le fichier brut augmenterait le risque.",
            ),
            (
                "Cartes simplifiées. ",
                "Seuil ± 2 écarts-types sur les moyennes de lots, pas les règles de Westgard ni les cibles constructeur. Suffisant pour montrer une dérive, pas pour remplacer un middleware d’analyseur.",
            ),
            (
                "Mono-poste. ",
                "Localhost, un fichier SQLite, pas de sauvegarde auto ni d’HTTPS. Le sujet était la traçabilité métier, pas l’exploitation cloud.",
            ),
            (
                "Comptes de démo. ",
                "Mot de passe connu (labo2026) pour le jury. Inacceptable en production, nécessaire pour une démo de groupe.",
            ),
            (
                "Audit et signalements incomplets. ",
                "On historise les champs du lot, pas « qui a cliqué ». Les signalements d’un lot sont recréés à chaque enregistrement : la liste est à jour, mais on perd le fil des allers-retours.",
            ),
            (
                "Peu de tests automatiques. ",
                "Le calendrier a privilégié les écrans visibles à l’oral. Une régression peut passer jusqu’à l’usage.",
            ),
        ],
    )
    add_para(
        doc,
        "Google OAuth et le code-barres du certificat restent illustratifs (secrets Google hors dépôt ; pas de lecteur physique).",
        italic=True,
        color=GRIS,
        space_after=8,
    )

    add_heading_custom(doc, "6. Pistes d’amélioration", 1)
    add_para(doc, "Pour chaque limite du § 5, une piste concrète, classée par facilité pour un groupe étudiant.")
    add_heading_custom(doc, "Priorité haute", 2)
    add_bullets(
        doc,
        [
            (
                "Tests métier. ",
                "Une vingtaine de cas sur les fonctions déjà isolées (colonne NIR → alerte ; lot sans machine → signalement ; validation > 12 mois → à revalider), lancés avant chaque démo.",
            ),
            (
                "Journal d’actions. ",
                "Une ligne utilisateur + action + lot + date à la validation, à la connexion et à l’acceptation de compte ; une page « Journal » suffit.",
            ),
            (
                "Signalements stables. ",
                "Ne plus les effacer : clôturer le motif disparu, en ouvrir un nouveau sinon, avec un commentaire facultatif.",
            ),
        ],
    )
    add_heading_custom(doc, "Priorité moyenne", 2)
    add_bullets(
        doc,
        [
            (
                "Cartes un peu plus réalistes. ",
                "Ajouter 2–3 règles de Westgard (ex. 1-3s, 2-2s) et dire laquelle a sauté.",
            ),
            (
                "Import CSV plus sûr. ",
                "Ne garder que les en-têtes et les valeurs numériques de contrôle ; détecter NIR / e-mail en mémoire, sans stocker le fichier brut.",
            ),
            (
                "Sauvegarde. ",
                "Bouton « exporter / importer la base » (c’est déjà un seul fichier).",
            ),
        ],
    )
    add_heading_custom(doc, "Plus tard, si le projet continue", 2)
    add_bullets(
        doc,
        [
            (
                "Déploiement. ",
                "Un poste partagé en salle plutôt que 15 copies OneDrive — seulement si un enseignant le demande.",
            ),
            (
                "Ne pas faire. ",
                "Un modèle d’IA sur de vrais bilans patients : ça sort du sujet et des obligations qu’on peut assumer.",
            ),
        ],
    )

    add_heading_custom(doc, "Pour la démo", 1)
    add_para(
        doc,
        "Ouvrir http://localhost:3000. Technicien : lot incomplet → signalement → correction. "
        "Biologiste : validation → certificat → synthèse PDF. Rappeler qu’aucune donnée patient réelle n’est utilisée. "
        "Le registre ne dit pas si un patient est malade : il dit si la donnée de calibration est complète, pseudonymisée, "
        "rattachée à une machine, et signée.",
    )

    out = Path(__file__).resolve().parent / "Rapport_Registre_Tracabilite_Laboratoires.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build()
